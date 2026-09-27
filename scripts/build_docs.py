"""Пересборка документации: docs/technical/openapi.json, code_reference.html и Word-версия документа.

openapi.json — спецификация FastAPI-приложения (то же, что отдает /openapi.json).
code_reference.html — справочник PyDoc по модулям app и ml_service в одном файле. Ссылки между
модулями ведут на якоря внутри страницы; локальные пути к исходникам заменяются путями от корня
репозитория. Скрипты из scripts/ в справочник не входят: они выполняют расчеты при импорте.

Запуск из корня проекта: python scripts/build_docs.py
"""
import datetime
import html
import importlib
import json
import pkgutil
import pydoc
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
OUT = ROOT / "docs" / "technical"
PACKAGES = ("app", "ml_service")


def modules() -> list:
    names = []
    for pkg_name in PACKAGES:
        pkg = importlib.import_module(pkg_name)
        names.append(pkg_name)
        for info in pkgutil.walk_packages(pkg.__path__, prefix=f"{pkg_name}."):
            names.append(info.name)
    return sorted(names)


def build_openapi() -> Path:
    from app.api import app
    path = OUT / "openapi.json"
    path.write_text(json.dumps(app.openapi(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def build_code_reference(names: list) -> Path:
    doc = pydoc.HTMLDoc()
    # значения выводятся целиком: иначе PyDoc обрезает середину длинного пути, и его нельзя заменить
    full = pydoc.HTMLRepr()
    full.maxstring = full.maxother = 2000
    doc.repr, doc.escape = full.repr, full.escape
    known = set(names)
    sections, toc = [], []
    for name in names:
        module = importlib.import_module(name)
        body = doc.docmodule(module)
        # ссылки на другие модули проекта -> якоря внутри страницы; на стандартную библиотеку — убрать
        body = re.sub(r'href="([\w.]+)\.html(#[^"]*)?"',
                      lambda m: f'href="#{m.group(1)}"' if m.group(1) in known else 'href="#"', body)
        body = body.replace('<a href=".">index</a><br>', "")
        # абсолютный путь к исходнику -> путь от корня репозитория
        source = Path(module.__file__).resolve().relative_to(ROOT).as_posix()
        body = re.sub(r'<a href="file:[^"]*">[^<]*</a>', html.escape(source), body)
        # значения констант-путей (ROOT, MODELS, ...) показывают каталог сборки — заменить на корень проекта
        for root in {str(ROOT), ROOT.as_posix()}:
            body = body.replace(html.escape(root), "&lt;корень проекта&gt;").replace(root, "&lt;корень проекта&gt;")
        sections.append(f'<section id="{name}">{body}</section>')
        toc.append(f'<li><a href="#{name}">{name}</a></li>')
    page = (
        '<!doctype html><html lang="ru"><meta charset="utf-8"><title>Справочник модулей</title>'
        "<style>body{font-family:Arial,sans-serif;max-width:1200px;margin:32px auto;padding:0 16px;"
        "line-height:1.4}table{max-width:100%}pre{white-space:pre-wrap}section{margin-top:40px}</style>"
        "<h1>Справочник модулей: app и ml_service</h1>"
        "<p>Сгенерировано PyDoc из исходного кода командой <code>python scripts/build_docs.py</code>.</p>"
        f"<ul>{''.join(toc)}</ul>{''.join(sections)}</html>\n"
    )
    path = OUT / "code_reference.html"
    path.write_text(page, encoding="utf-8")
    return path


_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _runs(par, text: str) -> None:
    """Строка markdown -> runs абзаца: **жирный**, `код`, [ссылка](адрес) -> текст ссылки."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("["):
            par.add_run(part[1:part.index("]")])
        else:
            par.add_run(part)


def build_docx(md_path: Path, out_path: Path) -> Path | None:
    """Word-версия технической документации из markdown (заголовки, абзацы, списки, таблицы, код, рисунки)."""
    try:
        from docx import Document
        from docx.shared import Cm, Pt
    except ImportError:
        print("python-docx не установлен: .docx не собран (pip install python-docx)")
        return None
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.5)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i, para = 0, []

    def flush() -> None:
        if para:
            _runs(doc.add_paragraph(), " ".join(para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            flush()
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                run = doc.add_paragraph().add_run(lines[i])
                run.font.name, run.font.size = "Consolas", Pt(9)
                i += 1
        elif m := re.match(r"(#{1,4}) (.*)", line):
            flush()
            level = len(m.group(1))
            doc.add_heading(m.group(2), level=0 if level == 1 else level - 1)
        elif m := re.match(r"!\[[^\]]*\]\(([^)]+)\)", line):
            flush()
            doc.add_picture(str(md_path.parent / m.group(1)), width=Cm(16))
        elif line.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, text in enumerate(cells[:len(rows[0])]):
                    cell_par = table.cell(r, c).paragraphs[0]
                    _runs(cell_par, text)
                    if r == 0:
                        for run in cell_par.runs:
                            run.bold = True
            continue
        elif m := re.match(r"(\s*)(- |\d+\. )(.*)", line):
            flush()
            style = "List Bullet" if m.group(2) == "- " else "List Number"
            _runs(doc.add_paragraph(style=style), m.group(3))
        elif not line.strip():
            flush()
        else:
            para.append(line.strip())
        i += 1
    flush()
    props = doc.core_properties   # шаблон python-docx подписывает документ собой — перезаписать
    props.author = props.last_modified_by = props.comments = props.keywords = props.subject = ""
    props.title = "Предиктор задержек наземного транспорта. Техническая документация"
    props.created = props.modified = datetime.datetime.now()
    doc.save(str(out_path))
    return out_path


def main() -> None:
    names = modules()
    built = [build_openapi(), build_code_reference(names),
             build_docx(OUT / "technical_documentation.md", OUT / "Техническая_документация_MVP.docx")]
    for path in filter(None, built):
        print(f"{path.relative_to(ROOT).as_posix()}: {path.stat().st_size / 1024:.0f} КБ")
    print(f"модулей в справочнике: {len(names)}")


if __name__ == "__main__":
    main()
